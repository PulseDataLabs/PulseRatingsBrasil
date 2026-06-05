#!/usr/bin/env python
# coding: utf-8
"""
Scraper: Moody's – Lista de Entidades Vigentes (Brasil)
Fonte:   https://moodyslocal.com.br/
Saída:   data/moodys_entidades.csv
"""
import os
import sys
import re
import datetime
from io import BytesIO

import pandas as pd
from bs4 import BeautifulSoup

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scrapers.utils.base import BaseScraper

BASE_URL = "https://moodyslocal.com.br"


class MoodysEntidadesScraper(BaseScraper):
    name = "moodys_entidades"
    group = "ratings"
    enabled = True
    phase = 1
    accumulate = False
    chaves_dedup = ["no_entidade"]

    # Catálogo de Metadados
    title = "Moody's — Emissores"
    description = "Lista de entidades (emissores corporativos, financeiros e públicos) com rating vigente pela Moody's no Brasil."
    icon = "M"
    icon_class = "icon-moodys"
    badge = "Diário"
    badge_class = "badge-daily"
    tags = ["ratings", "moodys", "entidades", "emissores"]
    source = "Moody's"

    def fetch(self) -> pd.DataFrame:
        from curl_cffi import requests
        from openpyxl import load_workbook

        session = requests.Session()

        self.logger.info(f"Acessando {BASE_URL} para buscar o link do Excel...")
        try:
            resp = session.get(BASE_URL, impersonate="chrome", timeout=60)
            resp.raise_for_status()
        except Exception as e:
            self.logger.error(f"Erro ao acessar Moody's Local: {e}")
            return pd.DataFrame()

        soup = BeautifulSoup(resp.text, "html.parser")
        download_url = None

        for a in soup.find_all("a", href=True):
            text = a.get_text(strip=True)
            href = a["href"]
            if "Lista de Classificações Vigentes" in text or re.search(r"MOODYS_LOCAL_BRAZIL.*\.xlsx?", href, re.IGNORECASE):
                download_url = href if href.startswith("http") else BASE_URL.rstrip("/") + "/" + href.lstrip("/")
                break

        if not download_url:
            self.logger.error("Link de download da Moody's Local não encontrado.")
            return pd.DataFrame()

        self.logger.info(f"Baixando arquivo Excel: {download_url}")
        try:
            resp_file = session.get(download_url, impersonate="chrome", timeout=120)
            resp_file.raise_for_status()
        except Exception as e:
            self.logger.error(f"Erro ao baixar o Excel: {e}")
            return pd.DataFrame()

        xlsx_bytes = BytesIO(resp_file.content)

        self.logger.info("Carregando planilha...")
        try:
            wb = load_workbook(filename=xlsx_bytes, read_only=True, data_only=True)
            sheet = wb.active
        except Exception as e:
            self.logger.error(f"Erro ao carregar planilha: {e}")
            return pd.DataFrame()

        self.logger.info("Processando linhas da planilha de entidades...")
        headers = None
        data_rows = []

        try:
            for row in sheet.iter_rows(values_only=True):
                # Busca linha do cabeçalho
                if headers is None:
                    if "Emissor" in row and "Rating / Avaliação" in row:
                        headers = [str(cell).strip() if cell is not None else "" for cell in row]
                        while headers and headers[-1] == "":
                            headers.pop()
                        continue
                else:
                    emissor_val = row[1] if len(row) > 1 else None
                    if emissor_val is None or str(emissor_val).strip() in ("", "None", "-"):
                        break

                    cleaned_row = list(row[:len(headers)])
                    if len(cleaned_row) < len(headers):
                        cleaned_row += [None] * (len(headers) - len(cleaned_row))
                    data_rows.append(cleaned_row)
        finally:
            wb.close()

        if not data_rows or not headers:
            return pd.DataFrame()

        df = pd.DataFrame(data_rows, columns=headers)
        df.rename(columns={"Emissor": "no_entidade"}, inplace=True)
        df.drop_duplicates(subset=["no_entidade"], inplace=True)
        
        today_str = datetime.date.today().strftime("%Y-%m-%d")
        df.insert(0, "dt_captura", today_str)
        df["link"] = ""

        # Mantém apenas as colunas unificadas
        df = df[["dt_captura", "no_entidade", "link"]]
        return df


if __name__ == "__main__":
    MoodysEntidadesScraper().run()
