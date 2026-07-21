#!/usr/bin/env python
# coding: utf-8
"""
Scraper: Moody's – Lista de Emissores Vigentes (Brasil)
Fonte:   https://moodyslocal.com.br/
Saída:   data/moodys_emissores.csv
"""
import os
import sys
import re
import datetime
from io import BytesIO

import pandas as pd
from bs4 import BeautifulSoup

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scrapers.utils.base import BaseScraper
from scripts.utils import print_done, print_info, print_warn, print_start, print_fail

BASE_URL = "https://moodyslocal.com.br"


class MoodysEmissoresScraper(BaseScraper):
    name = "moodys_emissores"
    group = "ratings"
    enabled = True
    phase = 1
    accumulate = False
    chaves_dedup = ["no_emissor"]

    # Catálogo de Metadados
    title = "Moody's — Emissores"
    description = "Lista de emissores (emissores corporativos, financeiros e públicos) com rating vigente pela Moody's no Brasil."
    icon = "M"
    icon_class = "icon-moodys"
    badge = "Diário"
    badge_class = "badge-daily"
    tags = ["ratings", "moodys", "emissores"]
    source = "Moody's"

    def fetch(self) -> pd.DataFrame:
        from openpyxl import load_workbook
        from scrapers.utils.moodys_helper import get_moodys_session

        self.logger.info(f"Acessando {BASE_URL} para buscar o link do Excel...")
        print_start("Acessando moodyslocal.com.br...")
        try:
            session, proxies = get_moodys_session(self.logger, BASE_URL)
        except Exception as e:
            self.logger.error(f"Erro ao obter sessão Moody's Local: {e}")
            print_fail(f"Erro ao obter sessão: {e}")
            return pd.DataFrame()

        try:
            resp = session.get(BASE_URL, impersonate="chrome", timeout=60)
            resp.raise_for_status()
        except Exception as e:
            self.logger.error(f"Erro ao acessar Moody's Local: {e}")
            print_fail(f"Erro ao acessar Moody's Local: {e}")
            return pd.DataFrame()

        soup = BeautifulSoup(resp.text, "html.parser")
        download_url = None

        for a in soup.find_all("a", href=True):
            text = a.get_text(strip=True)
            href = a["href"]
            if "Lista de Classificações Vigentes" in text or re.search(r"MOODYS_LOCAL_BRAZIL.*\.xlsx?", href, re.IGNORECASE):
                download_url = href if href.startswith("http") else BASE_URL.rstrip("/") + "/" + href.lstrip("/")
                break

        if not download_url:
            self.logger.error("Link de download da Moody's Local não encontrado.")
            print_fail("Link do Excel não encontrado na página")
            return pd.DataFrame()

        self.logger.info(f"Baixando arquivo Excel: {download_url}")
        print_start(f"Baixando Excel...")
        try:
            resp_file = session.get(download_url, impersonate="chrome", timeout=120)
            resp_file.raise_for_status()
        except Exception as e:
            self.logger.error(f"Erro ao baixar o Excel: {e}")
            print_fail(f"Erro ao baixar Excel: {e}")
            return pd.DataFrame()
        print_done("Excel baixado")

        xlsx_bytes = BytesIO(resp_file.content)

        self.logger.info("Carregando planilha...")
        print_start("Carregando planilha...")
        try:
            wb = load_workbook(filename=xlsx_bytes, read_only=True, data_only=True)
            sheet = wb.active
        except Exception as e:
            self.logger.error(f"Erro ao carregar planilha: {e}")
            print_fail(f"Erro ao carregar planilha: {e}")
            return pd.DataFrame()

        self.logger.info("Processando linhas da planilha de emissores...")
        print_start("Processando emissores...")
        headers = None
        data_rows = []

        try:
            for row in sheet.iter_rows(values_only=True):
                # Busca linha do cabeçalho
                if headers is None:
                    if "Emissor" in row and "Rating / Avaliação" in row:
                        headers = [str(cell).strip() if cell is not None else "" for cell in row]
                        while headers and headers[-1] == "":
                            headers.pop()
                        continue
                else:
                    emissor_val = row[1] if len(row) > 1 else None
                    if emissor_val is None or str(emissor_val).strip() in ("", "None", "-"):
                        break

                    cleaned_row = list(row[:len(headers)])
                    if len(cleaned_row) < len(headers):
                        cleaned_row += [None] * (len(headers) - len(cleaned_row))
                    data_rows.append(cleaned_row)
        finally:
            wb.close()

        if not data_rows or not headers:
            return pd.DataFrame()

        df = pd.DataFrame(data_rows, columns=headers)
        df.rename(columns={"Emissor": "no_emissor"}, inplace=True)
        df.drop_duplicates(subset=["no_emissor"], inplace=True)
        
        today_str = datetime.date.today().strftime("%Y-%m-%d")
        df.insert(0, "dt_captura", today_str)
        df["link"] = ""

        # Mantém apenas as colunas unificadas
        df = df[["dt_captura", "no_emissor", "link"]]

        print_done(f"{len(df)} emissores extraídos")
        return df


if __name__ == "__main__":
    MoodysEmissoresScraper().run()
